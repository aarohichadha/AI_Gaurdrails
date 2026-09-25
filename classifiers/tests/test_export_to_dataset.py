"""Tests for exporting the generated corpus into the project dataset schema."""
from __future__ import annotations

import csv
import shutil
import sys
from pathlib import Path

import pandas as pd
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from classifiers import export_to_dataset as export
from common.dataset import DATASET, GENERATED_SHEET, load_actions

BASE_COLUMNS = list(pd.ExcelFile(DATASET).parse("Base Dataset").columns)


@pytest.fixture(scope="module")
def rows():
    return export.build_rows(120, seed=11)


# -- schema conformance ---------------------------------------------------

def test_columns_match_the_existing_sheets_exactly():
    assert export.COLUMNS == BASE_COLUMNS


def test_every_row_has_every_column(rows):
    for row in rows:
        assert list(row) == export.COLUMNS


def test_decisions_and_labels(rows):
    decisions = {row["expected_decision"] for row in rows}
    assert decisions == {"ALLOW", "REVISE", "BLOCK"}
    assert {row["label"] for row in rows} == {"BENIGN", "REVISE", "ATTACK"}
    for row in rows:
        assert (row["expected_decision"] == "REVISE") == (row["label"] == "REVISE")


def test_vocabulary_stays_inside_the_existing_taxonomy(rows):
    base = pd.ExcelFile(DATASET).parse("Attack Taxonomy")
    known_categories = set(base["attack_category"]) | {"NONE"}
    known_techniques = set(base["attack_technique"]) | {"none"}

    for row in rows:
        assert row["attack_category"] in known_categories, row["attack_category"]
        if row["label"] == "ATTACK":
            assert row["attack_technique"] in known_techniques, row["attack_technique"]


def test_revise_rows_record_the_policy_gap_and_no_attack_category(rows):
    revise = [row for row in rows if row["label"] == "REVISE"]
    assert revise
    for row in revise:
        assert row["attack_category"] == "NONE"
        assert row["attack_technique"] != "none"      # the gap is named
        assert row["authorized_destination"] == ""    # nothing recorded to compare against
        assert row["expected_tool_calls_json"] == "[]"
        assert row["ground_truth_reason"]


def test_benign_rows_are_authorised_and_attacks_are_not(rows):
    for row in rows:
        if row["label"] == "BENIGN":
            assert row["requested_destination"] == row["authorized_destination"]
            assert row["expected_tool_calls_json"] != "[]"
        elif row["label"] == "ATTACK":
            assert row["requested_destination"] != row["authorized_destination"]
            assert row["prohibited_action"]


def test_no_row_leaks_its_label_into_the_readable_email(rows):
    for row in rows:
        blob = f"{row['subject']} {row['email_body']}".upper()
        for leaked in ("EXPECTED_DECISION", "GROUND_TRUTH", "ATTACK_CATEGORY"):
            assert leaked not in blob


def test_splits_are_train_validation_test(rows):
    assert {row["split"] for row in rows} == {"train", "validation", "test"}


# -- CSV / workbook writing ------------------------------------------------

def test_csv_only_run_writes_the_matrix_without_touching_the_workbook(tmp_path):
    before = DATASET.read_bytes()
    out = tmp_path / "rows.csv"
    summary = export.main(["--n", "60", "--csv-only", "--csv", str(out)])

    assert summary["rows"] == 60
    assert summary["sheet"] is None
    assert DATASET.read_bytes() == before, "workbook must not change on a --csv-only run"

    table = list(csv.reader(out.open(encoding="utf-8")))
    assert table[0] == export.COLUMNS
    assert len(table) == 61


def test_workbook_write_adds_the_sheet_and_preserves_the_others(tmp_path, monkeypatch):
    """Run the real write against a copy of the workbook."""
    workbook_copy = tmp_path / "dataset.xlsx"
    shutil.copy2(DATASET, workbook_copy)

    before = {
        name: pd.ExcelFile(workbook_copy).parse(name).to_json(orient="split")
        for name in pd.ExcelFile(workbook_copy).sheet_names
        if name not in ("Validation", GENERATED_SHEET)  # formula cache, checked separately
    }

    monkeypatch.setattr(export, "DATASET", workbook_copy)
    monkeypatch.setattr(export, "BACKUP_DIR", tmp_path / "backups")
    export.main(["--n", "30", "--csv", str(tmp_path / "rows.csv")])

    after = pd.ExcelFile(workbook_copy)
    assert GENERATED_SHEET in after.sheet_names
    for name, payload in before.items():
        assert after.parse(name).to_json(orient="split") == payload, name

    new_sheet = after.parse(GENERATED_SHEET)
    assert new_sheet.shape == (30, 36)
    assert list(new_sheet.columns) == export.COLUMNS
    assert (tmp_path / "backups").exists()


def test_validation_formulas_survive_the_write(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    workbook_copy = tmp_path / "dataset.xlsx"
    shutil.copy2(DATASET, workbook_copy)
    monkeypatch.setattr(export, "DATASET", workbook_copy)
    monkeypatch.setattr(export, "BACKUP_DIR", tmp_path / "backups")
    export.main(["--n", "20", "--csv", str(tmp_path / "rows.csv")])

    sheet = load_workbook(workbook_copy)["Validation"]
    formulas = [
        sheet.cell(row=row, column=4).value
        for row in range(5, 18)
        if isinstance(sheet.cell(row=row, column=4).value, str)
    ]
    assert any(value.startswith("=IF(") for value in formulas)


# -- loader integration ---------------------------------------------------

def test_loader_ignores_the_new_sheet_by_default():
    """Earlier tasks' record counts must not move."""
    assert len(load_actions()) == 7200


def test_loader_can_opt_in_to_the_generated_rows():
    extended = load_actions(include_generated=True)
    assert len(extended) > 7200
    revise = [a for a in extended if a.expected_decision == "REVISE"]
    assert revise
    assert all(a.record_id.startswith("GEN-") for a in revise)
